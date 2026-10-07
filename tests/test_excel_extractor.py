from openpyxl import Workbook

from ingestion.excel_extractor import (
    extract_sheet_records,
    find_kpi_sheet_header,
    filter_glossaries_for_retrieval,
    selected_worksheets,
)


def metadata() -> dict[str, object]:
    return {
        "original_name": "glossaire_kpis_test",
        "local_filename": "test.xlsx",
    }


def test_curly_apostrophe_kpi_header_is_recognized() -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "KPI"
    sheet.append(["KPI’s", "Définition", "Formule"])
    sheet.append(["Appels reçus", "Nombre d'appels", "SUM(appels)"])

    records = extract_sheet_records(sheet, metadata())

    assert len(records) == 1
    assert records[0]["kpi_name"] == "Appels reçus"


def test_unnamed_kpi_column_is_recovered_on_kpi_sheet() -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "KPIs"
    sheet.append(
        ["Scope", "Atelier", None, "Format", "Description", "Formule de calcul"]
    )
    sheet.append(
        ["Intelcia", "A1", "DMT", "00.00", "Durée moyenne", "AVG(x)", "ODCalls"]
    )
    sheet.append(["Intelcia", "A1", "Volume", "0", "Nombre", "COUNT(x)"])

    header = find_kpi_sheet_header(sheet)
    records = extract_sheet_records(sheet, metadata())

    assert header is not None
    assert header[1]["kpi_name"] == 2
    assert [record["kpi_name"] for record in records] == ["DMT", "Volume"]
    assert records[0]["_original_columns"]["Nom KPI"] == "DMT"
    assert records[0]["_original_columns"]["Unnamed column G"] == "ODCalls"


def test_auxiliary_recipe_sheet_is_ignored() -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "recette"
    sheet.append(["Nom KPI", "Commentaire"])
    sheet.append(["DMT", "OK"])

    assert extract_sheet_records(sheet, metadata()) == []


def test_kpis_sheet_is_preferred_over_configured_fallback(monkeypatch) -> None:
    workbook = Workbook()
    workbook.active.title = "Catalogue"
    workbook.create_sheet("KPIS")
    monkeypatch.setattr(
        "ingestion.excel_extractor.KPI_SHEETS", {"test": ("Catalogue",)}
    )

    selected = selected_worksheets(workbook, "test")

    assert [sheet.title for sheet in selected] == ["KPIS", "Catalogue"]


def test_configured_sheets_are_used_without_kpi_sheet(monkeypatch) -> None:
    workbook = Workbook()
    workbook.active.title = "Catalogue"
    workbook.create_sheet("Notes")
    monkeypatch.setattr(
        "ingestion.excel_extractor.KPI_SHEETS", {"test": ("Catalogue",)}
    )

    selected = selected_worksheets(workbook, "test")

    assert [sheet.title for sheet in selected] == ["Catalogue", "Notes"]


def test_apostrophe_kpi_sheet_is_recognized() -> None:
    workbook = Workbook()
    workbook.active.title = "KPI's"
    workbook.create_sheet("Autre feuille")

    selected = selected_worksheets(workbook, "cimr")

    assert [sheet.title for sheet in selected] == ["KPI's", "Autre feuille"]


def test_retrieval_output_uses_configured_fallback(monkeypatch) -> None:
    monkeypatch.setattr(
        "ingestion.excel_extractor.KPI_SHEETS", {"test": ("Catalogue",)}
    )
    result = filter_glossaries_for_retrieval([
        {
            "glossary_name": "test",
            "kpis": [
                {"kpi_name": "A", "sheet_name": "Catalogue"},
                {"kpi_name": "B", "sheet_name": "Notes"},
            ],
        }
    ])
    assert [row["kpi_name"] for row in result[0]["kpis"]] == ["A"]
